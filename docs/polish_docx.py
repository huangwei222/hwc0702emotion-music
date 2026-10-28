"""polish_docx.py —— 对 html-to-docx 产出的 .docx 做中文排版收尾。

html-to-docx 只映射了 CSS 里能直译的属性，有三件中文技术文档的惯例它做不到：
  1. 标题字体：转换器会把所有 CJK 文字的 eastAsia 字体统一成正文的「宋体」，
     导致标题也是宋体，层次感不足 → 这里把 Heading1/2/3 改成「黑体 + Arial 西文」。
  2. 首行缩进：CSS text-indent 不映射 → 按 2 字符（firstLineChars=200）补齐。
  3. 正文页码从 1 开始：封面单独一节，若不重编号，正文第一页会显示「2」。

另外给每张表的首行加上「跨页重复标题行」，15 张表里有几张会跨页。

用法：
    python polish_docx.py in.docx out.docx
"""
import sys

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

HEADING_FONT_EA = '黑体'
HEADING_FONT_ASCII = 'Arial'


def set_run_fonts(run, eastasia, ascii_font):
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.insert(0, rFonts)
    rFonts.set(qn('w:ascii'), ascii_font)
    rFonts.set(qn('w:hAnsi'), ascii_font)
    rFonts.set(qn('w:eastAsia'), eastasia)
    rFonts.set(qn('w:cs'), ascii_font)


def is_code_paragraph(p):
    return any(r.font.name and 'Consolas' in str(r.font.name) for r in p.runs)


def fix_heading_fonts(doc):
    n = 0
    for p in doc.paragraphs:
        if not p.style.name.startswith('Heading'):
            continue
        for r in p.runs:
            set_run_fonts(r, HEADING_FONT_EA, HEADING_FONT_ASCII)
            n += 1
    return n


def fix_first_line_indent(doc):
    n = 0
    for p in doc.paragraphs:
        if p.style.name != 'Normal':
            continue
        if is_code_paragraph(p):
            continue
        if p.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER:
            continue
        text = p.text.strip()
        # 过短的（标签行）与题注（表 x-x / 图 x-x）不缩进
        if len(text) < 24 or text.startswith(('表 ', '图 ', '表　')):
            continue
        pPr = p._p.get_or_add_pPr()
        ind = pPr.find(qn('w:ind'))
        if ind is None:
            ind = OxmlElement('w:ind')
            pPr.append(ind)
        ind.set(qn('w:firstLineChars'), '200')   # 2 字符，随字号自适应
        ind.set(qn('w:firstLine'), '420')        # 兜底：21pt（10.5pt × 2）
        n += 1
    return n


def fix_page_number_restart(doc, section_index=1):
    """正文节页码从 1 开始（封面节不编号）。"""
    if len(doc.sections) <= section_index:
        return False
    sectPr = doc.sections[section_index]._sectPr
    if sectPr.find(qn('w:pgNumType')) is not None:
        return False
    pg = OxmlElement('w:pgNumType')
    pg.set(qn('w:start'), '1')
    # schema 要求 pgNumType 排在 cols / docGrid 之前
    anchor = None
    for tag in ('w:cols', 'w:docGrid', 'w:formProt'):
        anchor = sectPr.find(qn(tag))
        if anchor is not None:
            break
    if anchor is not None:
        anchor.addprevious(pg)
    else:
        sectPr.append(pg)
    return True


def repeat_table_headers(doc):
    n = 0
    for t in doc.tables:
        if not t.rows:
            continue
        tr = t.rows[0]._tr
        trPr = tr.find(qn('w:trPr'))
        if trPr is None:
            trPr = OxmlElement('w:trPr')
            tr.insert(0, trPr)
        if trPr.find(qn('w:tblHeader')) is None:
            trPr.append(OxmlElement('w:tblHeader'))
            n += 1
        if trPr.find(qn('w:cantSplit')) is None:
            trPr.append(OxmlElement('w:cantSplit'))
    return n


def main(src, dst):
    doc = docx.Document(src)
    a = fix_heading_fonts(doc)
    b = fix_first_line_indent(doc)
    c = fix_page_number_restart(doc)
    d = repeat_table_headers(doc)
    doc.save(dst)
    print(f'标题字体已改黑体: {a} 个 run')
    print(f'首行缩进已补: {b} 段')
    print(f'正文页码重起: {"是" if c else "已是"}')
    print(f'表格标题行跨页重复: {d} 张')
    print('saved ->', dst)


if __name__ == '__main__':
    if len(sys.argv) < 3:
        print(__doc__)
        sys.exit(1)
    main(sys.argv[1], sys.argv[2])
